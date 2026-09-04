from datetime import date

import pytest
from openpyxl import Workbook

from feishu_rag.xlsx_reader import XlsxExtractionError, read_xlsx_sections


def test_xlsx_rows_are_searchable_and_keep_sheet_name(tmp_path) -> None:
    path = tmp_path / "各岗位标准UPPH.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "岗位标准"
    sheet.append(["序号", "工段", "工序", "UPH"])
    sheet.append([1, "切割", "Cell AOI2线开机", 450])
    sheet.append([None, None, None, None])
    workbook.save(path)

    sections = read_xlsx_sections(path)

    assert len(sections) == 1
    assert sections[0].section == "岗位标准"
    assert "工序: Cell AOI2线开机" in sections[0].text
    assert "UPH: 450" in sections[0].text


def test_merged_header_value_is_available_to_each_data_row(tmp_path) -> None:
    path = tmp_path / "工时.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells("A1:A2")
    sheet["A1"] = "工段"
    sheet["B1"] = "常规系列"
    sheet["B2"] = "工时"
    sheet.append(["切割", 99])
    workbook.save(path)

    text = "\n".join(section.text for section in read_xlsx_sections(path))

    assert "工段: 切割" in text
    assert "常规系列 / 工时: 99" in text


def test_empty_sheets_and_rows_are_ignored(tmp_path) -> None:
    path = tmp_path / "工时.xlsx"
    workbook = Workbook()
    workbook.active.title = "空表"
    data = workbook.create_sheet("有效数据")
    data.append(["工段", "工时"])
    data.append([None, None])
    data.append(["绑定", 90])
    workbook.save(path)

    sections = read_xlsx_sections(path)

    assert [section.section for section in sections] == ["有效数据"]
    assert "工段: 绑定" in sections[0].text
    assert "工时: 90" in sections[0].text


def test_values_are_rendered_without_excel_display_formatting(tmp_path) -> None:
    path = tmp_path / "数据.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["日期", "比率", "启用"])
    sheet.append([date(2026, 9, 4), 0.125, True])
    workbook.save(path)

    text = read_xlsx_sections(path)[0].text

    assert "日期: 2026-09-04" in text
    assert "比率: 0.125" in text
    assert "启用: true" in text


def test_headerless_sheet_uses_stable_column_names(tmp_path) -> None:
    path = tmp_path / "无表头.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([1, 450])
    sheet.append([2, 500])
    workbook.save(path)

    text = read_xlsx_sections(path)[0].text

    assert "列1: 1" in text
    assert "列2: 450" in text
    assert "列1: 2" in text
    assert "列2: 500" in text


def test_corrupt_or_empty_workbook_is_rejected(tmp_path) -> None:
    corrupt = tmp_path / "损坏.xlsx"
    corrupt.write_bytes(b"not-a-workbook")
    with pytest.raises(XlsxExtractionError, match="无法读取"):
        read_xlsx_sections(corrupt)

    empty = tmp_path / "空.xlsx"
    Workbook().save(empty)
    with pytest.raises(XlsxExtractionError, match="没有可索引数据"):
        read_xlsx_sections(empty)
