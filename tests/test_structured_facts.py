from openpyxl import Workbook

from feishu_rag.structured_facts import extract_structured_facts


def test_extracts_upph_with_forward_filled_process_stage(tmp_path) -> None:
    path = tmp_path / "各岗位标准UPPH.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "各岗位标准UPPH"
    sheet.append(["序号", "工段", "工序", "UPH"])
    sheet.append([1, "切割", "Cell AOI2线开机", 450])
    sheet.append([2, None, "切割扫码", 200])
    workbook.save(path)

    facts = extract_structured_facts(path, "ie:upph")

    assert [(fact.operation_name, fact.process_stage, fact.numeric_value) for fact in facts] == [
        ("Cell AOI2线开机", "切割", 450.0),
        ("切割扫码", "切割", 200.0),
    ]
    assert {fact.fact_type for fact in facts} == {"upph"}


def test_extracts_series_stage_hours_and_preserves_zero(tmp_path) -> None:
    path = tmp_path / "系列工时.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "总表"
    sheet.append(["工段", "工时", None])
    sheet.append([None, "常规系列", "氧化物系列"])
    sheet.append(["绑定", 90, 101])
    sheet.append(["贴合", 0, 99])
    workbook.save(path)

    facts = extract_structured_facts(path, "ie:hours")

    values = {
        (fact.series_name, fact.process_stage): fact.numeric_value for fact in facts
    }
    assert values[("氧化物系列", "绑定")] == 101.0
    assert values[("常规系列", "贴合")] == 0.0


def test_extracts_part_number_stage_and_total_hours(tmp_path) -> None:
    path = tmp_path / "料号工时.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "By料号"
    sheet.append(["序号", "系列", "料号", "工时", None, "合计"])
    sheet.append([None, None, None, "切割", "绑定", None])
    sheet.append([1, "常规系列", "055010D", 99, 90, 1026])
    workbook.save(path)

    facts = extract_structured_facts(path, "ie:parts")

    values = {fact.process_stage: fact.numeric_value for fact in facts}
    assert values == {"切割": 99.0, "绑定": 90.0, "合计": 1026.0}
    assert {fact.part_number for fact in facts} == {"055010D"}


def test_unknown_sheet_returns_no_structured_facts(tmp_path) -> None:
    path = tmp_path / "其他.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "说明"
    sheet.append(["标题", "内容"])
    sheet.append(["说明", "普通文本"])
    workbook.save(path)

    assert extract_structured_facts(path, "other") == []
