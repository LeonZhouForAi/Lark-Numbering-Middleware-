"""从受控 IE 工作簿结构提取精确数值事实。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .models import StructuredFact


def _text(value: Any) -> str:
    return "" if value is None else " ".join(str(value).split())


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def _upph_facts(sheet: Any, source_id: str) -> list[StructuredFact] | None:
    header_row = None
    headers: list[str] = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = [_text(value) for value in row]
        if "工序" in values and any(value.upper() in {"UPH", "UPPH"} for value in values):
            header_row = row_number
            headers = values
            break
    if header_row is None:
        return None
    operation_col = headers.index("工序")
    metric_col = next(
        index for index, value in enumerate(headers) if value.upper() in {"UPH", "UPPH"}
    )
    stage_col = headers.index("工段") if "工段" in headers else None
    stage = ""
    facts: list[StructuredFact] = []
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if stage_col is not None and _text(row[stage_col]):
            stage = _text(row[stage_col])
        operation = _text(row[operation_col])
        value = _number(row[metric_col])
        if operation and value is not None:
            facts.append(
                StructuredFact(
                    source_id,
                    sheet.title,
                    row_number,
                    "upph",
                    process_stage=stage,
                    operation_name=operation,
                    metric_name="UPPH",
                    numeric_value=value,
                )
            )
    return facts


def _series_hour_facts(sheet: Any, source_id: str) -> list[StructuredFact] | None:
    rows = list(sheet.iter_rows(values_only=True))
    for index in range(len(rows) - 1):
        first = [_text(value) for value in rows[index]]
        second = [_text(value) for value in rows[index + 1]]
        if "工段" not in first or not any("系列" in value for value in second):
            continue
        stage_col = first.index("工段")
        series_columns = [
            (column, value)
            for column, value in enumerate(second)
            if "系列" in value
        ]
        facts: list[StructuredFact] = []
        for row_number, row in enumerate(rows[index + 2 :], start=index + 3):
            stage = _text(row[stage_col])
            if not stage:
                continue
            for column, series in series_columns:
                value = _number(row[column])
                if value is not None:
                    facts.append(
                        StructuredFact(
                            source_id,
                            sheet.title,
                            row_number,
                            "series_hours",
                            series_name=series,
                            process_stage=stage,
                            metric_name="工时",
                            numeric_value=value,
                        )
                    )
        return facts
    return None


def _part_hour_facts(sheet: Any, source_id: str) -> list[StructuredFact] | None:
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 3:
        return None
    first = [_text(value) for value in rows[0]]
    second = [_text(value) for value in rows[1]]
    if "料号" not in first or "系列" not in first:
        return None
    part_col = first.index("料号")
    series_col = first.index("系列")
    stage_columns = []
    for column in range(len(first)):
        stage = second[column] or first[column]
        if column > part_col and stage and stage not in {"工时"}:
            stage_columns.append((column, stage))
    facts: list[StructuredFact] = []
    for row_number, row in enumerate(rows[2:], start=3):
        part_number = _text(row[part_col])
        series = _text(row[series_col])
        if not part_number:
            continue
        for column, stage in stage_columns:
            value = _number(row[column])
            if value is not None:
                facts.append(
                    StructuredFact(
                        source_id,
                        sheet.title,
                        row_number,
                        "part_hours",
                        part_number=part_number,
                        series_name=series,
                        process_stage=stage,
                        metric_name="工时",
                        numeric_value=value,
                    )
                )
    return facts


def extract_structured_facts(path: Path, source_id: str) -> list[StructuredFact]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        facts: list[StructuredFact] = []
        for sheet in workbook.worksheets:
            extracted = (
                _upph_facts(sheet, source_id)
                or _part_hour_facts(sheet, source_id)
                or _series_hour_facts(sheet, source_id)
            )
            if extracted:
                facts.extend(extracted)
        return facts
    finally:
        workbook.close()
