"""本地文档解析和索引入口。"""

from __future__ import annotations

import argparse
import hashlib
import logging
import os
import re
from dataclasses import dataclass
from pathlib import Path

from .chunker import chunk_text
from .document_metadata import (
    DocumentMetadata,
    extract_document_metadata,
    resolve_lifecycle_states,
)
from .logging_utils import configure_logging
from .store import IndexStore, PreparedDocument
from .structured_facts import extract_structured_facts
from .semantic_chunker import AtomicUnit, SemanticPlanner, semantic_chunks
from .xlsx_reader import XlsxExtractionError, read_xlsx_sections


logger = logging.getLogger(__name__)


class UnsupportedFileError(ValueError):
    """文件格式需要先转换。"""


class DocumentExtractionError(RuntimeError):
    """文档内容无法完整抽取。"""


@dataclass(frozen=True)
class Section:
    text: str
    page: int | None = None
    section: str | None = None


SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".pdf", ".docx", ".xlsx"}


def _ocr_pdf_page(path: Path, page_no: int) -> str | None:
    """识别 PDF 指定页；依赖、渲染或识别失败时显式报错。"""
    try:
        import pytesseract
        from pdf2image import convert_from_path
    except ImportError as exc:
        raise DocumentExtractionError("PDF OCR 依赖不可用") from exc
    try:
        images = convert_from_path(str(path), first_page=page_no, last_page=page_no, dpi=220)
    except Exception as exc:
        raise DocumentExtractionError(f"PDF 第 {page_no} 页渲染失败") from exc
    if not images:
        raise DocumentExtractionError(f"PDF 第 {page_no} 页渲染失败")
    try:
        text = pytesseract.image_to_string(images[0], lang="chi_sim")
    except Exception:
        try:
            text = pytesseract.image_to_string(images[0], lang="eng")
        except Exception as exc:
            raise DocumentExtractionError(f"PDF 第 {page_no} 页 OCR 识别失败") from exc
    text = text.strip()
    return text or None


def _read_pdf(path: Path, enable_ocr: bool = True) -> list[Section]:
    from pypdf import PdfReader

    sections: list[Section] = []
    reader = PdfReader(str(path))
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(Section(text=text, page=page_number))
        elif enable_ocr:
            ocr_text = _ocr_pdf_page(path, page_number)
            if ocr_text:
                sections.append(Section(text=ocr_text, page=page_number))
    return sections


def _read_docx(path: Path) -> list[Section]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))

    def table_rows(table: Table, seen_cells: set[object]) -> list[str]:
        rows: list[str] = []
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                cell_element = cell._tc
                if cell_element in seen_cells:
                    continue
                seen_cells.add(cell_element)
                cell_text = ""
                for cell_content in cell.iter_inner_content():
                    if isinstance(cell_content, Paragraph):
                        text = cell_content.text.strip()
                        if text:
                            cell_text = f"{cell_text} {text}".strip()
                    elif isinstance(cell_content, Table):
                        nested_rows = table_rows(cell_content, seen_cells)
                        if nested_rows:
                            nested_text = " / ".join(nested_rows)
                            cell_text = f"{cell_text} / {nested_text}".strip(" / ")
                cells.append(cell_text)
            if any(cells):
                rows.append(" | ".join(cells))
        return rows

    parts: list[str] = []
    for content in document.iter_inner_content():
        if isinstance(content, Paragraph):
            text = content.text.strip()
            if text:
                parts.append(text)
        elif isinstance(content, Table):
            parts.extend(table_rows(content, set()))
    text = "\n\n".join(parts)
    return [Section(text=text)] if text else []


def extract_sections(path: str | Path, enable_ocr: bool = True) -> list[Section]:
    """按格式抽取文本；`.doc` 和 `.wps` 明确要求先转 DOCX。"""

    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix in {".doc", ".wps"}:
        raise UnsupportedFileError(f"{file_path.name} 先转换为 DOCX 后再索引")
    if suffix in {".txt", ".md", ".markdown"}:
        text = file_path.read_text(encoding="utf-8")
        return [Section(text=text)] if text.strip() else []
    if suffix == ".pdf":
        return _read_pdf(file_path, enable_ocr=enable_ocr)
    if suffix == ".docx":
        return _read_docx(file_path)
    if suffix == ".xlsx":
        try:
            xlsx_sections = read_xlsx_sections(file_path)
        except XlsxExtractionError as exc:
            raise DocumentExtractionError(str(exc)) from exc
        return [
            Section(text=section.text, section=section.section)
            for section in xlsx_sections
        ]
    raise UnsupportedFileError(f"不支持的文件格式: {file_path.suffix or '(无扩展名)'}")


def _file_checksum(
    path: Path,
    enable_ocr: bool,
    chunk_strategy_version: str,
    chunk_model: str,
    parser_version: str,
) -> str:
    content_checksum = hashlib.sha256(path.read_bytes()).hexdigest()
    ocr_cache_mode = str(enable_ocr).lower() if path.suffix.lower() == ".pdf" else "na"
    return hashlib.sha256(
        f"{content_checksum}:{chunk_strategy_version}:{chunk_model}:"
        f"parser={parser_version}:ocr={ocr_cache_mode}".encode("utf-8")
    ).hexdigest()


def _local_source_id(root: Path, relative_path: str) -> str:
    root_hash = hashlib.sha256(str(root.resolve()).encode("utf-8")).hexdigest()
    return f"local:{root_hash}:{relative_path}"


def _prepare_file(
    path: Path,
    root: Path,
    max_chars: int = 900,
    enable_ocr: bool = True,
    semantic_planner: SemanticPlanner | None = None,
    chunk_strategy_version: str = "local-v1",
    chunk_model: str = "",
    parser_version: str = "parser-v2",
    metadata: DocumentMetadata | None = None,
) -> PreparedDocument | None:
    relative_path = path.relative_to(root).as_posix()
    source_id = _local_source_id(root, relative_path)
    title = path.stem
    checksum = _file_checksum(
        path, enable_ocr, chunk_strategy_version, chunk_model, parser_version
    )
    sections = extract_sections(path, enable_ocr=enable_ocr)
    overlap = max(0, min(120, max_chars // 5))

    def local_chunks() -> list:
        chunks = []
        for section in sections:
            chunks.extend(
                chunk_text(
                    section.text,
                    source_id=source_id,
                    title=title,
                    max_chars=max_chars,
                    overlap=overlap,
                    page=section.page,
                    section=section.section,
                )
            )
        return chunks

    if semantic_planner is None:
        chunks = local_chunks()
    else:
        try:
            units = []
            for section in sections:
                paragraphs = [part.strip() for part in re.split(r"\n\s*\n+", section.text) if part.strip()]
                for paragraph in paragraphs:
                    units.append(AtomicUnit(f"{source_id}:unit-{len(units)}", paragraph, section.page, section.section))
            chunks = semantic_chunks(units, source_id, title, semantic_planner, max_chars)
        except Exception as exc:
            logger.warning(
                "semantic_chunk_fallback source_id=%s error_type=%s",
                source_id,
                type(exc).__name__,
            )
            chunks = local_chunks()
    if not chunks:
        return None
    document_metadata = metadata or extract_document_metadata(path.name)
    structured_facts = (
        tuple(extract_structured_facts(path, source_id))
        if path.suffix.lower() == ".xlsx"
        else ()
    )
    return PreparedDocument(
        source_id,
        title,
        relative_path,
        checksum,
        tuple(chunks),
        document_code=document_metadata.document_code or "",
        document_version=document_metadata.document_version or "",
        effective_date=document_metadata.effective_date or "",
        lifecycle_state=document_metadata.lifecycle_state,
        parser_version=parser_version,
        decision_reason=document_metadata.decision_reason,
        structured_facts=structured_facts,
    )


def index_file(
    path: Path,
    root: Path,
    store: IndexStore,
    max_chars: int = 900,
    enable_ocr: bool = True,
    semantic_planner: SemanticPlanner | None = None,
    chunk_strategy_version: str = "local-v1",
    chunk_model: str = "",
    parser_version: str = "parser-v2",
) -> bool:
    source_id = _local_source_id(root, path.relative_to(root).as_posix())
    checksum = _file_checksum(
        path, enable_ocr, chunk_strategy_version, chunk_model, parser_version
    )
    if store.document_checksum(source_id) == checksum:
        return False
    prepared = _prepare_file(
        path,
        root,
        max_chars=max_chars,
        enable_ocr=enable_ocr,
        semantic_planner=semantic_planner,
        chunk_strategy_version=chunk_strategy_version,
        chunk_model=chunk_model,
        parser_version=parser_version,
        metadata=extract_document_metadata(path.name),
    )
    if prepared is None:
        return False
    updated, _ = store.apply_document_snapshot((prepared,))
    return bool(updated)


def index_directory(
    root: str | Path,
    store: IndexStore,
    max_chars: int = 900,
    enable_ocr: bool = True,
    semantic_planner: SemanticPlanner | None = None,
    chunk_strategy_version: str = "local-v1",
    chunk_model: str = "",
    parser_version: str = "parser-v2",
) -> int:
    """递归索引目录，返回成功索引的文件数量。"""

    root_path = Path(root).resolve()
    if not root_path.is_dir():
        raise FileNotFoundError(f"local root is not a directory: {root_path}")
    prepared_updates: list[PreparedDocument] = []
    retained_source_ids: set[str] = set()
    supported_paths = [
        path
        for path in sorted(root_path.rglob("*"))
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    ]
    metadata_by_path = resolve_lifecycle_states(
        {
            path.relative_to(root_path).as_posix(): extract_document_metadata(path.name)
            for path in supported_paths
        }
    )
    for path in supported_paths:
        relative_path = path.relative_to(root_path).as_posix()
        metadata = metadata_by_path[relative_path]
        source_id = _local_source_id(root_path, relative_path)
        checksum = _file_checksum(
            path, enable_ocr, chunk_strategy_version, chunk_model, parser_version
        )
        if store.document_checksum(source_id) == checksum:
            prepared_updates.append(
                PreparedDocument(
                    source_id,
                    path.stem,
                    relative_path,
                    checksum,
                    None,
                    document_code=metadata.document_code or "",
                    document_version=metadata.document_version or "",
                    effective_date=metadata.effective_date or "",
                    lifecycle_state=metadata.lifecycle_state,
                    parser_version=parser_version,
                    decision_reason=metadata.decision_reason,
                )
            )
            retained_source_ids.add(source_id)
            continue
        prepared = _prepare_file(
            path,
            root_path,
            max_chars=max_chars,
            enable_ocr=enable_ocr,
            semantic_planner=semantic_planner,
            chunk_strategy_version=chunk_strategy_version,
            chunk_model=chunk_model,
            parser_version=parser_version,
            metadata=metadata,
        )
        if prepared is not None:
            prepared_updates.append(prepared)
            retained_source_ids.add(prepared.source_id)
    indexed, _ = store.apply_document_snapshot(
        prepared_updates,
        local_root=root_path,
        retained=retained_source_ids,
    )
    return indexed


def main() -> None:
    parser = argparse.ArgumentParser(
        description="索引本地 PDF/DOCX/XLSX/Markdown/TXT 文件"
    )
    parser.add_argument("root", type=Path, help="文档目录")
    parser.add_argument("--db", type=Path, default=Path("./data/rag.sqlite3"), help="SQLite 数据库路径")
    parser.add_argument("--max-chars", type=int, default=900)
    parser.add_argument("--no-ocr", action="store_true", help="禁用扫描 PDF OCR")
    args = parser.parse_args()
    configure_logging(os.getenv("LOG_LEVEL", "INFO"))
    store = IndexStore(args.db)
    try:
        count = index_directory(
            args.root,
            store,
            max_chars=args.max_chars,
            enable_ocr=not args.no_ocr,
            parser_version=os.getenv("RAG_PARSER_VERSION", "parser-v2").strip()
            or "parser-v2",
        )
        print(f"indexed_files={count} indexed_chunks={store.count_chunks()}")
    finally:
        store.close()


if __name__ == "__main__":
    main()
