"""将 XLSX 工作表转换为稳定、可检索的文本行。"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any


class XlsxExtractionError(RuntimeError):
    """工作簿无法被完整、可靠地抽取。"""


@dataclass(frozen=True)
class XlsxSection:
    text: str
    section: str


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).replace("\n", " ").split())


def _all_strings(values: list[Any]) -> bool:
    populated = [value for value in values if value not in (None, "")]
    return bool(populated) and all(isinstance(value, str) for value in populated)


def _sheet_rows(sheet: Any) -> tuple[list[list[Any]], list[list[Any]], bool]:
    max_row = int(sheet.max_row or 0)
    max_column = int(sheet.max_column or 0)
    if max_row < 1 or max_column < 1:
        return [], [], False

    original = [
        [sheet.cell(row=row, column=column).value for column in range(1, max_column + 1)]
        for row in range(1, max_row + 1)
    ]
    projected = [row[:] for row in original]
    header_merge = False
    for merged_range in sheet.merged_cells.ranges:
        anchor = original[merged_range.min_row - 1][merged_range.min_col - 1]
        if merged_range.min_row <= 3:
            header_merge = True
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                projected[row - 1][column - 1] = anchor
    return original, projected, header_merge


def _render_sheet(sheet: Any) -> XlsxSection | None:
    original, projected, header_merge = _sheet_rows(sheet)
    populated_indexes = [
        index for index, row in enumerate(original) if any(value not in (None, "") for value in row)
    ]
    if not populated_indexes:
        return None

    first_index = populated_indexes[0]
    first_row = original[first_index]
    header_count = 1 if _all_strings(first_row) else 0
    first_has_blanks = any(value in (None, "") for value in first_row)
    allow_stacked_header = header_count == 1 and (first_has_blanks or header_merge)
    if allow_stacked_header:
        for index in populated_indexes[1:3]:
            if index != first_index + header_count or not _all_strings(original[index]):
                break
            header_count += 1

    width = len(first_row)
    if header_count:
        headers: list[str] = []
        for column in range(width):
            parts: list[str] = []
            for row in projected[first_index : first_index + header_count]:
                value = _text(row[column])
                if value and value not in parts:
                    parts.append(value)
            headers.append(" / ".join(parts) or f"列{column + 1}")
        data_start = first_index + header_count
    else:
        headers = [f"列{column + 1}" for column in range(width)]
        data_start = first_index

    rendered_rows: list[str] = []
    for index in range(data_start, len(original)):
        if not any(value not in (None, "") for value in original[index]):
            continue
        fields = [
            f"{headers[column]}: {_text(projected[index][column])}"
            for column in range(width)
            if _text(projected[index][column])
        ]
        if fields:
            rendered_rows.append(" | ".join(fields))
    if not rendered_rows:
        return None
    return XlsxSection("\n\n".join(rendered_rows), str(sheet.title))


def read_xlsx_sections(path: Path) -> list[XlsxSection]:
    """Read cached cell values without modifying the source workbook."""

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=False, data_only=True)
    except Exception as exc:
        raise XlsxExtractionError(f"无法读取 XLSX：{path.name}") from exc
    try:
        sections = [
            section
            for sheet in workbook.worksheets
            if (section := _render_sheet(sheet)) is not None
        ]
    finally:
        workbook.close()
    if not sections:
        raise XlsxExtractionError(f"XLSX 没有可索引数据：{path.name}")
    return sections
